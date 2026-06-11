#!/usr/bin/env python3
"""Audit all legacy competition data sources and emit a migration report."""

from __future__ import annotations

import gzip
import json
import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
LEGACY = ROOT.parent / "旧程序数据"
OUT = ROOT / "public" / "legacy-source-audit.json"

SQL_SOURCES = {
    "software_cloud": LEGACY / "software_cloud.sql",
    "software_config": LEGACY / "software_config.sql",
    "api_vmp": LEGACY / "api_vmp_20231031_190745.sql.gz",
    "api_vmp_20240121": LEGACY / "db_api_vmp_20240121_032801_vuydtz.sql.gz",
}

EXCEL_SOURCES = {
    "competition_users": LEGACY / "产赛用户.xlsx",
    "judges": LEGACY / "评委数据.xlsx",
}

MINIPROGRAM = LEGACY / "小程序mp-vote"


def read_text(path: Path) -> str:
    if path.suffix == ".gz":
        with gzip.open(path, "rt", errors="ignore") as file:
            return file.read()
    return path.read_text(errors="ignore")


def parse_columns(create_body: str) -> list[dict[str, str]]:
    columns = []
    for raw_line in create_body.splitlines():
        line = raw_line.strip().rstrip(",")
        if not line.startswith("`"):
            continue
        match = re.match(r"`(?P<name>[^`]+)`\s+(?P<type>[^\s,]+)(?P<rest>.*)", line)
        if not match:
            continue
        columns.append(
            {
                "name": match.group("name"),
                "type": match.group("type"),
                "raw": line,
            }
        )
    return columns


def count_insert_rows(values_sql: str) -> int:
    rows = 0
    depth = 0
    in_quote = False
    escaped = False
    for char in values_sql:
        if escaped:
            escaped = False
            continue
        if in_quote and char == "\\":
            escaped = True
            continue
        if char == "'":
            in_quote = not in_quote
            continue
        if not in_quote and char == "(":
            if depth == 0:
                rows += 1
            depth += 1
        elif not in_quote and char == ")":
            depth -= 1
    return rows


def audit_sql(name: str, path: Path) -> dict[str, Any]:
    text = read_text(path)
    tables: dict[str, Any] = {}
    for match in re.finditer(r"CREATE TABLE `(?P<table>[^`]+)` \((?P<body>.*?)\)\s+ENGINE=", text, re.S):
        table = match.group("table")
        tables[table] = {
            "columns": parse_columns(match.group("body")),
            "insertRows": 0,
        }
    for match in re.finditer(r"INSERT INTO `(?P<table>[^`]+)`(?: \((?P<cols>.*?)\))? VALUES (?P<vals>.*?);", text, re.S):
        table = match.group("table")
        tables.setdefault(table, {"columns": [], "insertRows": 0})
        tables[table]["insertRows"] += count_insert_rows(match.group("vals"))
    return {
        "source": name,
        "file": str(path.relative_to(ROOT.parent)),
        "tableCount": len(tables),
        "tables": tables,
    }


def audit_excel(name: str, path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.worksheets[0]
    rows = list(sheet.iter_rows(values_only=True))
    headers = [str(value).strip() if value is not None else f"列{index + 1}" for index, value in enumerate(rows[0])]
    non_empty = 0
    for row in rows[1:]:
        if any(value not in (None, "") for value in row):
            non_empty += 1
    return {
        "source": name,
        "file": str(path.relative_to(ROOT.parent)),
        "sheet": sheet.title,
        "rows": non_empty,
        "columns": headers,
    }


def audit_miniprogram() -> dict[str, Any]:
    js_files = sorted(MINIPROGRAM.glob("pages/**/*.js")) + [MINIPROGRAM / "app.js"]
    endpoints: dict[str, set[str]] = {}
    pages = []
    for path in js_files:
        if not path.exists():
            continue
        rel = str(path.relative_to(MINIPROGRAM))
        text = path.read_text(errors="ignore")
        pages.append(rel)
        for endpoint in re.findall(r"url:\s*['\"]([^'\"]+)['\"]", text):
            if endpoint.startswith("/pages/"):
                continue
            endpoints.setdefault(endpoint, set()).add(rel)
    return {
        "source": "mp-vote",
        "file": str(MINIPROGRAM.relative_to(ROOT.parent)),
        "pageCount": len(pages),
        "pages": pages,
        "endpointCount": len(endpoints),
        "endpoints": {key: sorted(value) for key, value in sorted(endpoints.items())},
    }


def main() -> None:
    payload = {
        "sql": {name: audit_sql(name, path) for name, path in SQL_SOURCES.items()},
        "excel": {name: audit_excel(name, path) for name, path in EXCEL_SOURCES.items()},
        "miniprogram": audit_miniprogram(),
    }
    OUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({
        "sqlSources": {name: {"tables": value["tableCount"]} for name, value in payload["sql"].items()},
        "excelSources": {name: {"rows": value["rows"], "columns": len(value["columns"])} for name, value in payload["excel"].items()},
        "miniprogramEndpoints": payload["miniprogram"]["endpointCount"],
        "output": str(OUT.relative_to(ROOT)),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
