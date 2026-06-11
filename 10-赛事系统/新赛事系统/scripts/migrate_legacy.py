#!/usr/bin/env python3
"""Build a read-only JSON snapshot from the legacy competition exports."""

from __future__ import annotations

import json
import gzip
import re
from collections import Counter, defaultdict
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
LEGACY = ROOT.parent / "旧程序数据"
SQL_FILE = LEGACY / "software_cloud.sql"
API_VMP_SQL = LEGACY / "db_api_vmp_20240121_032801_vuydtz.sql.gz"
USER_XLSX = LEGACY / "产赛用户.xlsx"
JUDGE_XLSX = LEGACY / "评委数据.xlsx"
OUT = ROOT / "public" / "legacy-data.js"
DEDUP_REPORT = ROOT / "public" / "legacy-dedupe-report.json"

TABLES = [
    "dfwq_competition",
    "dfwq_competition_scorerules",
    "dfwq_comps_user",
    "dfwq_group",
    "dfwq_groupjudges",
    "dfwq_integral",
    "dfwq_player",
    "dfwq_score",
    "sys_user",
]

API_VMP_OFFSET = {
    "competitions": 100000,
    "groups": 200000,
    "players": 300000,
    "users": 400000,
    "judges": 500000,
}


def split_sql_values(raw: str) -> list[Any]:
    values: list[str | None] = []
    current: list[str] = []
    in_quote = False
    escaped = False
    was_quoted = False

    for char in raw:
        if escaped:
            current.append(
                {
                    "n": "\n",
                    "r": "\r",
                    "t": "\t",
                    "0": "\0",
                    "'": "'",
                    "\\": "\\",
                }.get(char, char)
            )
            escaped = False
            continue
        if in_quote and char == "\\":
            escaped = True
            continue
        if char == "'":
            in_quote = not in_quote
            was_quoted = True
            continue
        if char == "," and not in_quote:
            values.append(convert_value("".join(current), was_quoted))
            current = []
            was_quoted = False
            continue
        current.append(char)

    values.append(convert_value("".join(current), was_quoted))
    return values


def split_insert_rows(raw: str) -> list[list[Any]]:
    rows: list[list[Any]] = []
    current: list[str] = []
    depth = 0
    in_quote = False
    escaped = False

    for char in raw:
        if escaped:
            current.append("\\" + char)
            escaped = False
            continue
        if in_quote and char == "\\":
            escaped = True
            continue
        if char == "'":
            in_quote = not in_quote
            current.append(char)
            continue
        if not in_quote and char == "(":
            if depth == 0:
                current = []
            else:
                current.append(char)
            depth += 1
            continue
        if not in_quote and char == ")":
            depth -= 1
            if depth == 0:
                rows.append(split_sql_values("".join(current)))
            else:
                current.append(char)
            continue
        if depth > 0:
            current.append(char)

    return rows


def convert_value(value: str, quoted: bool) -> Any:
    value = value.strip()
    if not quoted and value.upper() == "NULL":
        return None
    if quoted:
        return value
    if re.fullmatch(r"-?\d+", value):
        try:
            return int(value)
        except ValueError:
            return value
    if re.fullmatch(r"-?\d+\.\d+", value):
        try:
            return float(value)
        except ValueError:
            return value
    return value


def read_sql(path: Path) -> str:
    if path.suffix == ".gz":
        with gzip.open(path, "rt", errors="ignore") as file:
            return file.read()
    return path.read_text(errors="ignore")


def parse_sql_columns(sql: str) -> dict[str, list[str]]:
    schemas: dict[str, list[str]] = {}
    for match in re.finditer(r"CREATE TABLE `(?P<table>[^`]+)` \((?P<body>.*?)\)\s+ENGINE=", sql, re.S):
        columns = []
        for line in match.group("body").splitlines():
            line = line.strip()
            if line.startswith("`"):
                columns.append(line.split("`", 2)[1])
        schemas[match.group("table")] = columns
    return schemas


def parse_sql(sql: str, include_tables: list[str] | None = None) -> dict[str, list[dict[str, Any]]]:
    schemas = parse_sql_columns(sql)
    tables = include_tables or sorted(schemas)
    data: dict[str, list[dict[str, Any]]] = {table: [] for table in tables}
    pattern = re.compile(
        r"INSERT INTO `(?P<table>[^`]+)`(?: \((?P<cols>.*?)\))? VALUES (?P<vals>.*?);",
        re.S,
    )
    for match in pattern.finditer(sql):
        table = match.group("table")
        if table not in data:
            continue
        if match.group("cols"):
            columns = [part.strip().strip("`") for part in match.group("cols").split(",")]
        else:
            columns = schemas.get(table, [])
        for values in split_insert_rows(match.group("vals")):
            data[table].append(dict(zip(columns, values)))
    return data


def read_sheet(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value).strip() if value else f"列{idx + 1}" for idx, value in enumerate(next(rows))]
    records = []
    for row in rows:
        record = {headers[idx]: normalize_excel_value(value) for idx, value in enumerate(row)}
        if any(value not in (None, "") for value in record.values()):
            records.append(record)
    return records


def normalize_excel_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    return value


def index_by(rows: list[dict[str, Any]], key: str) -> dict[Any, dict[str, Any]]:
    return {row.get(key): row for row in rows if row.get(key) is not None}


def ts(value: Any) -> str | None:
    if value in (None, "", 0, "0"):
        return None
    try:
        number = int(str(value))
    except (TypeError, ValueError):
        return str(value)
    return datetime.fromtimestamp(number).strftime("%Y-%m-%d %H:%M:%S")


def offset(value: Any, key: str) -> int | None:
    if value in (None, ""):
        return None
    try:
        return API_VMP_OFFSET[key] + int(value)
    except (TypeError, ValueError):
        return None


def compact_key(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "").lower())


def phone_key(value: Any) -> str:
    return re.sub(r"\D+", "", str(value or ""))


def day_key(value: Any) -> str:
    return str(value or "")[:10]


def latest_text(*values: Any) -> str:
    return max([str(value or "") for value in values], default="")


def prefer_value(current: Any, candidate: Any) -> Any:
    if current in (None, "", 0, "0") and candidate not in (None, ""):
        return candidate
    return current


def merge_records(base: dict[str, Any], extra: dict[str, Any]) -> dict[str, Any]:
    merged = dict(base)
    for key, value in extra.items():
        if key not in merged:
            merged[key] = value
        else:
            merged[key] = prefer_value(merged[key], value)
    sources = []
    for row in (base, extra):
        source = row.get("legacy_source") or row.get("source")
        legacy_id = row.get("legacy_id") or row.get("id")
        if source or legacy_id:
            sources.append({"source": source, "legacyId": legacy_id})
    if sources:
        merged.setdefault("merged_sources", [])
        existing = {(item.get("source"), item.get("legacyId")) for item in merged["merged_sources"]}
        for item in sources:
            marker = (item.get("source"), item.get("legacyId"))
            if marker not in existing:
                merged["merged_sources"].append(item)
                existing.add(marker)
    return merged


def best_row(rows: list[dict[str, Any]], kind: str) -> dict[str, Any]:
    def score(row: dict[str, Any]) -> tuple[Any, ...]:
        common = (
            1 if row.get("del") in (None, 0, "0") else 0,
            1 if row.get("banned") in (None, 0, "0") else 0,
        )
        if kind == "competition":
            return common + (
                1 if row.get("signup") in (1, "1") else 0,
                row.get("playerCount") or 0,
                latest_text(row.get("startDate"), row.get("comps_time")),
                row.get("id") or 0,
            )
        if kind == "group":
            return common + (
                1 if row.get("end") in (None, 0, "0") else 0,
                -(row.get("sort") or 0),
                row.get("id") or 0,
            )
        if kind == "player":
            return common + (
                1 if row.get("isPaySuccess") in (1, "1") else 0,
                1 if row.get("complete") in (1, "1") else 0,
                1 if Number(row.get("score")) > 0 else 0,
                latest_text(row.get("paid_time"), row.get("player_date")),
                row.get("id") or 0,
            )
        if kind == "score":
            return (
                1 if Number(row.get("score")) > 0 else 0,
                Number(row.get("score")),
                latest_text(row.get("score_time"), row.get("date")),
                row.get("id") or 0,
            )
        if kind == "judge":
            return common + (
                1 if phone_key(row.get("tel")) else 0,
                row.get("id") or 0,
            )
        if kind == "user":
            return common + (
                1 if row.get("openid") else 0,
                latest_text(row.get("user_date"), row.get("active")),
                row.get("id") or 0,
            )
        return common + (row.get("id") or 0,)

    return max(rows, key=score)


def Number(value: Any) -> float:
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def merge_duplicate_rows(rows: list[dict[str, Any]], key_fn, kind: str) -> tuple[list[dict[str, Any]], dict[Any, Any], dict[str, Any]]:
    buckets: dict[Any, list[dict[str, Any]]] = defaultdict(list)
    passthrough = []
    for row in rows:
        key = key_fn(row)
        key_parts = key if isinstance(key, tuple) else (key,)
        if all(part not in (None, "") for part in key_parts):
            buckets[key].append(row)
        else:
            passthrough.append(row)

    merged_rows = list(passthrough)
    id_map = {}
    duplicate_groups = []
    duplicate_rows = 0
    for key, bucket in buckets.items():
        if len(bucket) == 1:
            merged_rows.append(bucket[0])
            continue
        keep = best_row(bucket, kind)
        merged = dict(keep)
        for item in bucket:
            if item is keep:
                continue
            merged = merge_records(merged, item)
        merged_rows.append(merged)
        duplicate_groups.append({"key": str(key), "keptId": keep.get("id"), "mergedIds": [item.get("id") for item in bucket if item.get("id") != keep.get("id")]})
        duplicate_rows += len(bucket) - 1
        for item in bucket:
            if item.get("id") is not None and keep.get("id") is not None:
                id_map[item.get("id")] = keep.get("id")

    return merged_rows, id_map, {"groups": len(duplicate_groups), "rowsMerged": duplicate_rows, "examples": duplicate_groups[:20]}


def remap_field(rows: list[dict[str, Any]], field: str, id_map: dict[Any, Any]) -> int:
    changed = 0
    for row in rows:
        value = row.get(field)
        if value in id_map and id_map[value] != value:
            row[field] = id_map[value]
            changed += 1
    return changed


def recompute_competition_summaries(data: dict[str, Any]) -> None:
    players_by_comp = defaultdict(list)
    groups_by_comp = defaultdict(list)
    scores_by_comp = Counter()
    judges_by_comp: dict[Any, set[Any]] = defaultdict(set)
    for player in data["players"]:
        players_by_comp[player.get("cid")].append(player)
    for group in data["groups"]:
        groups_by_comp[group.get("comps_id")].append(group)
    for score in data["scores"]:
        scores_by_comp[score.get("comps_id")] += 1
        judges_by_comp[score.get("comps_id")].add(score.get("judges_id"))
    for comp in data["competitions"]:
        cid = comp.get("id")
        comp_players = players_by_comp.get(cid, [])
        comp["groupCount"] = len(groups_by_comp.get(cid, []))
        comp["playerCount"] = len(comp_players)
        comp["paidPlayerCount"] = sum(1 for item in comp_players if item.get("isPaySuccess") == 1)
        comp["scoredPlayerCount"] = sum(1 for item in comp_players if Number(item.get("score")) > 0)
        comp["judgeCount"] = len(judges_by_comp.get(cid, set()))
        comp["scoreRecordCount"] = scores_by_comp.get(cid, 0)


def dedupe_normalized(data: dict[str, Any]) -> tuple[dict[str, Any], dict[str, Any]]:
    before = {key: len(data[key]) for key in ("competitions", "groups", "users", "judges", "players", "scores")}
    report: dict[str, Any] = {"before": before, "steps": {}, "remappedReferences": {}}

    data["competitions"], comp_map, report["steps"]["competitions"] = merge_duplicate_rows(
        data["competitions"],
        lambda row: (compact_key(row.get("comps_name")), day_key(row.get("startDate"))),
        "competition",
    )
    report["remappedReferences"]["groups.comps_id"] = remap_field(data["groups"], "comps_id", comp_map)
    report["remappedReferences"]["players.cid"] = remap_field(data["players"], "cid", comp_map)
    report["remappedReferences"]["scores.comps_id"] = remap_field(data["scores"], "comps_id", comp_map)

    data["groups"], group_map, report["steps"]["groups"] = merge_duplicate_rows(
        data["groups"],
        lambda row: (row.get("comps_id"), compact_key(row.get("name") or row.get("group_name"))),
        "group",
    )
    report["remappedReferences"]["players.tid"] = remap_field(data["players"], "tid", group_map)
    report["remappedReferences"]["scores.group_id"] = remap_field(data["scores"], "group_id", group_map)

    data["users"], user_map, report["steps"]["users"] = merge_duplicate_rows(
        data["users"],
        lambda row: row.get("openid") or row.get("unionid"),
        "user",
    )
    report["remappedReferences"]["players.uid"] = remap_field(data["players"], "uid", user_map)
    report["remappedReferences"]["judges.uid"] = remap_field(data["judges"], "uid", user_map)

    data["judges"], judge_map, report["steps"]["judges"] = merge_duplicate_rows(
        data["judges"],
        lambda row: phone_key(row.get("tel")),
        "judge",
    )
    report["remappedReferences"]["scores.judges_id"] = remap_field(data["scores"], "judges_id", judge_map)

    data["players"], player_map, report["steps"]["players"] = merge_duplicate_rows(
        data["players"],
        lambda row: (
            row.get("cid"),
            row.get("tid"),
            compact_key(row.get("name")),
            phone_key(row.get("mobile")),
            compact_key(row.get("show_name")),
        ),
        "player",
    )
    report["remappedReferences"]["scores.player_id"] = remap_field(data["scores"], "player_id", player_map)

    data["scores"], _, report["steps"]["scores"] = merge_duplicate_rows(
        data["scores"],
        lambda row: (row.get("judges_id"), row.get("player_id")),
        "score",
    )

    recompute_competition_summaries(data)
    report["after"] = {key: len(data[key]) for key in ("competitions", "groups", "users", "judges", "players", "scores")}
    report["removed"] = {key: report["before"][key] - report["after"][key] for key in report["before"]}
    data["dedupe"] = report
    return data, report


def prune_invalid_competitions(data: dict[str, Any]) -> dict[str, Any]:
    groups_by_comp = defaultdict(list)
    players_by_comp = defaultdict(list)
    scores_by_comp = defaultdict(list)
    for group in data["groups"]:
        groups_by_comp[group.get("comps_id")].append(group)
    for player in data["players"]:
        players_by_comp[player.get("cid")].append(player)
    for score in data["scores"]:
        scores_by_comp[score.get("comps_id")].append(score)

    valid_competition_ids = {
        comp.get("id")
        for comp in data["competitions"]
        if groups_by_comp.get(comp.get("id")) and players_by_comp.get(comp.get("id")) and scores_by_comp.get(comp.get("id"))
    }
    invalid_competitions = [comp for comp in data["competitions"] if comp.get("id") not in valid_competition_ids]

    before = {key: len(data[key]) for key in ("competitions", "groups", "players", "scores")}
    data["competitions"] = [comp for comp in data["competitions"] if comp.get("id") in valid_competition_ids]
    data["groups"] = [group for group in data["groups"] if group.get("comps_id") in valid_competition_ids]
    data["players"] = [player for player in data["players"] if player.get("cid") in valid_competition_ids]
    data["scores"] = [score for score in data["scores"] if score.get("comps_id") in valid_competition_ids]
    recompute_competition_summaries(data)

    report = {
        "rule": "赛事必须同时具备组别、选手/报名和评分记录，否则从 normalized 业务视图剔除。",
        "before": before,
        "after": {key: len(data[key]) for key in before},
        "removed": {key: before[key] - len(data[key]) for key in before},
        "invalidCompetitions": [
            {
                "id": comp.get("id"),
                "name": comp.get("comps_name"),
                "groupCount": len(groups_by_comp.get(comp.get("id"), [])),
                "playerCount": len(players_by_comp.get(comp.get("id"), [])),
                "scoreRecordCount": len(scores_by_comp.get(comp.get("id"), [])),
            }
            for comp in invalid_competitions
        ],
    }
    data["invalidCompetitionPrune"] = report
    return report


def build_api_vmp_normalized(api: dict[str, list[dict[str, Any]]]) -> dict[str, list[dict[str, Any]]]:
    signups_by_id = index_by(api.get("vmpdb_signups", []), "sid")
    shows_by_signup = index_by(api.get("vmpdb_shows", []), "signupID")
    groups_by_id = index_by(api.get("vmpdb_types", []), "tid")
    voters_by_id = index_by(api.get("vmpdb_voters", []), "vid")

    competitions = []
    for comp in api.get("vmpdb_comps", []):
        competitions.append(
            {
                "id": offset(comp.get("cid"), "competitions"),
                "legacy_source": "api_vmp_20240121",
                "legacy_id": comp.get("cid"),
                "comps_name": comp.get("name"),
                "startDate": ts(comp.get("startDate")),
                "endDate": ts(comp.get("endDate")),
                "sort": comp.get("sort"),
                "signup": comp.get("signup"),
                "comps_time": ts(comp.get("date")),
                "tel": comp.get("tel"),
                "citys": comp.get("location"),
                "wxChannelId": comp.get("wxChannelId"),
                "vidLogo": comp.get("vidLogo"),
                "vidQrcode": comp.get("vidQrcode"),
                "videoUrl": comp.get("wxChannelVid"),
                "comps_posterPic": comp.get("img"),
                "comps_defaultAvatar": comp.get("avatar"),
                "intro": comp.get("intro"),
                "sign_price": 0,
                "create_by": None,
                "player_num": 0,
                "member_visibility": 0,
                "del": comp.get("del"),
                "end": comp.get("end"),
                "uname": comp.get("uname"),
            }
        )

    groups = []
    for group in api.get("vmpdb_types", []):
        groups.append(
            {
                "id": offset(group.get("tid"), "groups"),
                "legacy_source": "api_vmp_20240121",
                "legacy_id": group.get("tid"),
                "comps_id": offset(group.get("cid"), "competitions"),
                "name": group.get("name"),
                "group_name": group.get("name"),
                "sort": group.get("sort"),
                "cost": group.get("cost"),
                "del": group.get("del"),
                "end": group.get("end"),
                "voters": group.get("voters"),
            }
        )

    users = []
    for user in api.get("vmpdb_users", []):
        users.append(
            {
                "id": offset(user.get("uid"), "users"),
                "legacy_source": "api_vmp_20240121",
                "legacy_id": user.get("uid"),
                "openid": user.get("openid"),
                "unionid": user.get("unionid"),
                "nickname": user.get("nickname"),
                "avatarUrl": user.get("avatarUrl"),
                "sex": user.get("sex"),
                "user_date": ts(user.get("date")),
                "active": ts(user.get("active")),
                "banned": user.get("banned"),
                "is_voter": user.get("is_voter"),
                "rankScore": user.get("rankScore"),
            }
        )

    judges = []
    for voter in api.get("vmpdb_voters", []):
        user = next((item for item in users if item.get("legacy_id") == voter.get("uid")), None)
        judges.append(
            {
                "id": offset(voter.get("vid"), "judges"),
                "legacy_source": "api_vmp_20240121",
                "legacy_id": voter.get("vid"),
                "nickname": voter.get("name") or user.get("nickname") if user else voter.get("name"),
                "tel": voter.get("mobile"),
                "uid": offset(voter.get("uid"), "users"),
                "source": "api_vmp",
                "user_state": 0 if voter.get("del") == 0 else 1,
            }
        )

    players = []
    for signup in api.get("vmpdb_signups", []):
        show = shows_by_signup.get(signup.get("sid"), {})
        players.append(
            {
                "id": offset(signup.get("sid"), "players"),
                "legacy_source": "api_vmp_20240121",
                "legacy_id": signup.get("sid"),
                "uid": offset(signup.get("uid"), "users"),
                "name": signup.get("name"),
                "show_name": signup.get("show_name") or show.get("name"),
                "tid": offset(signup.get("tid"), "groups"),
                "cid": offset(signup.get("cid"), "competitions"),
                "amount": signup.get("amount"),
                "cost": signup.get("cost"),
                "sn": signup.get("sn"),
                "mobile": signup.get("mobile"),
                "province": signup.get("province"),
                "city": signup.get("city"),
                "group_name": signup.get("group_name"),
                "zhidao": signup.get("zhidao"),
                "bianwu": signup.get("bianwu"),
                "style": signup.get("style"),
                "mp3_file": signup.get("file"),
                "player_date": ts(signup.get("date")),
                "paid_time": ts(signup.get("paid")),
                "isPaySuccess": 1 if signup.get("paid") else 0,
                "score": show.get("rating"),
                "sort": show.get("sort"),
                "complete": show.get("complete"),
                "del": signup.get("del"),
                "avatarUrl": show.get("avatar"),
                "costScore": 0,
            }
        )

    scores = []
    for index, vote in enumerate(api.get("vmpdb_votes", []), start=1):
        signup = signups_by_id.get(vote.get("sid"), {})
        group = groups_by_id.get(signup.get("tid"), {})
        scores.append(
            {
                "id": API_VMP_OFFSET["players"] + index,
                "legacy_source": "api_vmp_20240121",
                "judges_id": offset(vote.get("vid"), "judges"),
                "player_id": offset(vote.get("sid"), "players"),
                "group_id": offset(signup.get("tid"), "groups"),
                "comps_id": offset(signup.get("cid") or group.get("cid"), "competitions"),
                "score": vote.get("rating"),
                "detail1": vote.get("detail1"),
                "detail2": vote.get("detail2"),
                "detail3": vote.get("detail3"),
                "detail4": vote.get("detail4"),
                "detail5": vote.get("detail5"),
                "score_time": ts(vote.get("date")),
            }
        )

    return {
        "competitions": competitions,
        "groups": groups,
        "users": users,
        "judges": judges,
        "players": players,
        "scores": scores,
    }


def build_normalized(raw: dict[str, list[dict[str, Any]]], api: dict[str, list[dict[str, Any]]], excel_users: list[dict[str, Any]], excel_judges: list[dict[str, Any]]) -> dict[str, Any]:
    users_by_id = index_by(raw["dfwq_comps_user"], "id")
    groups_by_id = index_by(raw["dfwq_group"], "id")
    competitions = raw["dfwq_competition"]
    players = raw["dfwq_player"]
    scores = raw["dfwq_score"]
    judges_links = raw["dfwq_groupjudges"]

    players_by_comp = defaultdict(list)
    groups_by_comp = defaultdict(list)
    scores_by_comp = Counter()
    judges_by_comp: dict[Any, set[Any]] = defaultdict(set)
    for player in players:
        players_by_comp[player.get("cid")].append(player)
    for group in raw["dfwq_group"]:
        groups_by_comp[group.get("comps_id")].append(group)
    for score in scores:
        scores_by_comp[score.get("comps_id")] += 1
    for link in judges_links:
        judges_by_comp[link.get("comps_id")].add(link.get("judges_id"))

    api_normalized = build_api_vmp_normalized(api)

    comp_summaries = []
    for comp in competitions + api_normalized["competitions"]:
        cid = comp.get("id")
        comp_players = players_by_comp.get(cid, [])
        if comp.get("legacy_source") == "api_vmp_20240121":
            comp_players = [item for item in api_normalized["players"] if item.get("cid") == cid]
        paid = sum(1 for item in comp_players if item.get("isPaySuccess") == 1)
        scored = sum(1 for item in comp_players if item.get("score") not in (None, "", 0))
        comp_summaries.append(
            {
                **comp,
                "groupCount": len(groups_by_comp.get(cid, [])),
                "playerCount": len(comp_players),
                "paidPlayerCount": paid,
                "scoredPlayerCount": scored,
                "judgeCount": len(judges_by_comp.get(cid, set())),
                "scoreRecordCount": scores_by_comp.get(cid, 0),
            }
        )

    excel_user_ids = {row.get("用户ID") for row in excel_users}
    excel_judge_phones = {str(row.get("手机号码")).strip() for row in excel_judges if row.get("手机号码")}
    sql_judges = [user for user in raw["dfwq_comps_user"] if user.get("isJudges") == 1]
    merged_judges = []
    seen_judge_keys = set()
    for user in sql_judges:
        key = ("sql", user.get("id"))
        seen_judge_keys.add(key)
        if user.get("tel"):
            seen_judge_keys.add(("phone", str(user.get("tel")).strip()))
        merged_judges.append({**user, "source": "sql"})
    for row in excel_judges:
        phone = str(row.get("手机号码") or "").strip()
        if not phone:
            continue
        linked = next((user for user in raw["dfwq_comps_user"] if str(user.get("tel") or "").strip() == phone), None)
        key = ("phone", phone)
        if key in seen_judge_keys:
            continue
        seen_judge_keys.add(key)
        merged_judges.append(
            {
                "id": linked.get("id") if linked else None,
                "nickname": row.get("昵称"),
                "tel": phone,
                "source": "excel+sql" if linked else "excel",
                "user_state": linked.get("user_state") if linked else None,
            }
        )

    orphan_excel_users = [row for row in excel_users if row.get("用户ID") not in users_by_id]
    merged_judges.extend(api_normalized["judges"])

    return {
        "competitions": comp_summaries,
        "users": raw["dfwq_comps_user"] + api_normalized["users"],
        "excelUsers": excel_users,
        "judges": merged_judges,
        "groups": raw["dfwq_group"] + api_normalized["groups"],
        "groupJudges": judges_links,
        "players": players + api_normalized["players"],
        "scores": scores + api_normalized["scores"],
        "scoreRules": raw["dfwq_competition_scorerules"],
        "integral": raw["dfwq_integral"],
        "systemUsers": raw["sys_user"],
        "apiVmp": api_normalized,
        "checks": {
            "excelUserRows": len(excel_users),
            "excelUsersMatchedById": len(excel_user_ids & set(users_by_id.keys())),
            "excelUsersOnly": len(orphan_excel_users),
            "excelJudgeRows": len(excel_judges),
            "sqlJudgeRows": len(sql_judges),
            "mergedJudgeRows": len(merged_judges),
            "excelJudgePhones": len(excel_judge_phones),
            "apiVmpRows": sum(len(rows) for rows in api.values()),
            "apiVmpCompetitions": len(api_normalized["competitions"]),
            "apiVmpGroups": len(api_normalized["groups"]),
            "apiVmpPlayers": len(api_normalized["players"]),
            "apiVmpScores": len(api_normalized["scores"]),
        },
    }


def main() -> None:
    raw = parse_sql(read_sql(SQL_FILE), TABLES)
    api = parse_sql(read_sql(API_VMP_SQL))
    excel_users = read_sheet(USER_XLSX)
    excel_judges = read_sheet(JUDGE_XLSX)
    normalized, dedupe_report = dedupe_normalized(build_normalized(raw, api, excel_users, excel_judges))
    invalid_competition_report = prune_invalid_competitions(normalized)
    payload = {
        "meta": {
            "generatedAt": datetime.now().isoformat(timespec="seconds"),
            "sourceFiles": [str(SQL_FILE.name), str(API_VMP_SQL.name), str(USER_XLSX.name), str(JUDGE_XLSX.name)],
            "tableCounts": {table: len(rows) for table, rows in raw.items()},
            "apiVmpTableCounts": {table: len(rows) for table, rows in api.items()},
        },
        "raw": raw,
        "apiVmpRaw": api,
        "normalized": normalized,
    }
    OUT.write_text(
        "window.LEGACY_DATA = " + json.dumps(payload, ensure_ascii=False, separators=(",", ":")) + ";\n",
        encoding="utf-8",
    )
    DEDUP_REPORT.write_text(json.dumps(dedupe_report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(payload["meta"], ensure_ascii=False, indent=2))
    print(json.dumps(payload["normalized"]["checks"], ensure_ascii=False, indent=2))
    print(json.dumps({"dedupe": dedupe_report["removed"], "output": str(DEDUP_REPORT.relative_to(ROOT))}, ensure_ascii=False, indent=2))
    print(json.dumps({"invalidCompetitionsRemoved": invalid_competition_report["removed"]}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
